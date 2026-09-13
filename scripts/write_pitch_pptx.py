#!/usr/bin/env python3
"""Build an editable Vellisys client pitch deck (.pptx)."""

from __future__ import annotations

import os
import subprocess
import sys

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

NAVY = RGBColor(0x08, 0x14, 0x3A)
BLUE = RGBColor(0x1E, 0x4E, 0xFF)
GOLD = RGBColor(0xC6, 0xA1, 0x5B)
INK = RGBColor(0x14, 0x17, 0x12)
MUTED = RGBColor(0x5A, 0x61, 0x72)
PAPER = RGBColor(0xF4, 0xF6, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
MIST = RGBColor(0xC9, 0xD2, 0xEE)
SLIDE_W = Inches(13.333)
SLIDE_H = Inches(7.5)
A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"


def set_run(run, size, bold=False, color=INK, font="Arial"):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def add_box(slide, x, y, w, h, text, size=14, bold=False, color=INK, align=PP_ALIGN.LEFT, font="Arial", anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.auto_size = None
    try:
        tf._txBody.bodyPr.set("anchor", {MSO_ANCHOR.TOP: "t", MSO_ANCHOR.MIDDLE: "ctr", MSO_ANCHOR.BOTTOM: "b"}.get(anchor, "t"))
    except Exception:
        pass
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    set_run(run, size, bold, color, font)
    return box


def add_lines(slide, x, y, w, h, lines, size=13, color=INK, bold=False, spacing=1.08):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(6)
        run = p.add_run()
        run.text = line
        set_run(run, size, bold, color)
    return box


def fill_shape(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color
    shape.line.fill.background()


def add_rect(slide, x, y, w, h, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, x, y, w, h)
    fill_shape(sh, color)
    return sh


def set_picture_alpha(shape, percent: float) -> None:
    """percent is 0–100 visible opacity."""
    blip = shape._element.find(".//" + qn("a:blip"))
    if blip is None:
        return
    for child in list(blip):
        if child.tag == qn("a:alphaModFix"):
            blip.remove(child)
    amt = str(int(max(0, min(100, percent)) * 1000))
    el = etree.SubElement(blip, qn("a:alphaModFix"))
    el.set("amt", amt)


def set_slide_bg(slide, path: str) -> None:
    slide.shapes.add_picture(path, 0, 0, SLIDE_W, SLIDE_H)


def bars(slide) -> None:
    add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), BLUE)
    add_rect(slide, 0, SLIDE_H - Inches(0.08), SLIDE_W, Inches(0.08), BLUE)


def watermark(slide, mark, dark: bool) -> None:
    size = Inches(4.6)
    pic = slide.shapes.add_picture(
        mark,
        (SLIDE_W - size) / 2,
        (SLIDE_H - size) / 2 - Inches(0.1),
        size,
        size,
    )
    set_picture_alpha(pic, 18 if dark else 12)


def logo(slide, wordmark, dark: bool) -> None:
    h = Inches(0.42) if dark else Inches(0.36)
    w = h * (2154 / 633)
    x = Inches(0.7) if dark else SLIDE_W - Inches(0.7) - w
    y = Inches(0.22)
    slide.shapes.add_picture(wordmark, x, y, w, h)


def dark_rails(slide) -> None:
    add_rect(slide, 0, Inches(0.85), Inches(0.18), SLIDE_H - Inches(1.7), BLUE)
    add_rect(slide, Inches(0.18), Inches(0.85), Inches(0.07), SLIDE_H - Inches(1.7), GOLD)


def card(slide, x, y, w, h, fill=PAPER, accent=True):
    sh = add_rect(slide, x, y, w, h, fill)
    if accent:
        add_rect(slide, x, y, Inches(0.07), h, BLUE)
    return sh


def prepare_assets() -> str:
    root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    php = os.path.join(root, "scripts", "export_pitch_pptx_assets.php")
    out = subprocess.check_output(["php", php], text=True).strip().splitlines()[-1]
    return out


def build(asset_dir: str, dests: list[str]) -> None:
    mark = os.path.join(asset_dir, "mark-alpha.png")
    word = os.path.join(asset_dir, "wordmark-alpha.png")
    bg_dark = os.path.join(asset_dir, "bg-dark.png")
    bg_light = os.path.join(asset_dir, "bg-light.png")

    prs = Presentation()
    prs.slide_width = SLIDE_W
    prs.slide_height = SLIDE_H
    blank = prs.slide_layouts[6]

    def dark_slide():
        s = prs.slides.add_slide(blank)
        set_slide_bg(s, bg_dark)
        bars(s)
        watermark(s, mark, True)
        logo(s, word, True)
        return s

    def light_slide():
        s = prs.slides.add_slide(blank)
        set_slide_bg(s, bg_light)
        bars(s)
        watermark(s, mark, False)
        logo(s, word, False)
        return s

    # 1 Cover
    s = dark_slide()
    dark_rails(s)
    add_box(s, Inches(0.7), Inches(1.55), Inches(8), Inches(0.35), "CLIENT PITCH DECK", 14, True, GOLD)
    add_box(s, Inches(0.7), Inches(2.05), Inches(11), Inches(1.4), "Branded books for companies\nthat cannot lose the paper.", 32, True, WHITE)
    add_box(
        s,
        Inches(0.7),
        Inches(3.6),
        Inches(10.2),
        Inches(1.5),
        "Quotations, invoices, receipts, expenses, delivery notes, headed letters and reports in your logo, colours, currency and tax. One desk. Yearly packages. Built for East Africa and used anywhere a company still sends paper that must look like it left their office.",
        16,
        False,
        MIST,
    )
    add_box(s, Inches(0.7), Inches(5.3), Inches(11), Inches(0.35), "www.vellisys.com  ·  info@vellisys.com", 16, True, GOLD)
    add_box(s, Inches(0.7), Inches(5.65), Inches(11), Inches(0.3), "+256 779 971 024  ·  +256 756 524 451", 16, False, WHITE)
    add_box(s, Inches(0.7), Inches(6.0), Inches(11), Inches(0.3), "Front Star Digital  ·  Confidential to the company you are visiting", 13, False, MIST)
    add_box(s, Inches(0.7), Inches(6.55), Inches(11), Inches(0.4), "Send this pack before a demo. It is the whole system, the packages, and how we start together.", 13, False, WHITE)

    # 2 Pains
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.72), Inches(10), Inches(0.45), "The books are leaking.", 28, True, NAVY)
    add_box(s, Inches(0.55), Inches(1.18), Inches(10), Inches(0.35), "Common pains Vellisys is built to close", 16, False, MUTED)
    pains_l = [
        "Quotes live in Word. Invoices in Excel. Receipts in a pad. Nobody can show the same story twice.",
        "The sheet that leaves the office does not look like the company. Logo missing, colours wrong, tax guessed.",
        "Part payments are rewritten on the invoice. Debtors become a rumour.",
        "Staff email from personal Gmail. Clients never know if the bill is real.",
    ]
    pains_r = [
        "QuickBooks and Xero assume a monthly SaaS seat, US or UK payroll, and a bookkeeper who already lives in that product.",
        "Spreadsheets do not age, do not remind, and do not convert a quote to an invoice.",
        "Tax is hardcoded in the owner's head: 18% here, 16% there, until a sheet is wrong.",
        "Onboarding a new country means a new tool, a new login culture, and another invoice from Silicon Valley.",
    ]
    for col, items in enumerate((pains_l, pains_r)):
        for i, item in enumerate(items):
            x = Inches(0.55 + col * 6.35)
            y = Inches(1.65 + i * 1.35)
            card(s, x, y, Inches(6.1), Inches(1.22))
            add_box(s, x + Inches(0.22), y + Inches(0.12), Inches(5.7), Inches(1.0), item, 14, False, NAVY)

    # 3 What it is
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.72), Inches(10), Inches(0.45), "What Vellisys is", 28, True, NAVY)
    add_box(
        s,
        Inches(0.55),
        Inches(1.22),
        Inches(12.2),
        Inches(1.15),
        "Vellisys is branded books software. A company desk: the people who raise quotes, issue invoices, take receipts, log expenses, chase debtors and write headed letters. Super admin at Vellisys onboard the company, assign a sending mailbox, and keep the paid term honest.",
        16,
        False,
        INK,
    )
    bits = [
        ("YOUR PAPER", "Logo, two brand colours, one document design for the whole desk. Print, PDF, WhatsApp link or email from the company mailbox."),
        ("YOUR MONEY", "Any three-letter currency. USD on a sheet if you need it. You set the rate. Reports add it back to home currency."),
        ("YOUR TAX", "Name it VAT, GST, SST, IVA. Set the percent. Taxed lines use that rate. Old sheets keep the rate they were saved with."),
        ("YOUR SEATS", "One, two or three logins. Company admin plus Books or Sales access. Not an unlimited cloud that bills per extra user forever."),
    ]
    for i, (title, body) in enumerate(bits):
        x = Inches(0.55 + (i % 2) * 6.35)
        y = Inches(2.55 + (i // 2) * 2.2)
        card(s, x, y, Inches(6.1), Inches(2.0), PAPER, False)
        add_box(s, x + Inches(0.25), y + Inches(0.2), Inches(5.6), Inches(0.35), title, 14, True, BLUE)
        add_box(s, x + Inches(0.25), y + Inches(0.6), Inches(5.6), Inches(1.2), body, 15, False, INK)

    # 4 Compare
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.4), "Beside QuickBooks, Xero and Sage", 26, True, NAVY)
    add_box(s, Inches(0.55), Inches(1.12), Inches(12), Inches(0.3), "Different job. Different price. Different paper.", 15, False, MUTED)
    headers = ["", "Vellisys", "QuickBooks", "Xero", "Sage"]
    rows = [
        ["What it is", "Branded books desk", "Full SME accounting", "Cloud ledgers + apps", "Accounting suites"],
        ["Paper that leaves", "Your logo, colours, tax", "Generic / templates", "Generic / add-ons", "Forms, not stationery"],
        ["Billing", "Per year, 3 packages", "Monthly SaaS seats", "Monthly + payroll", "Licence + partners"],
        ["East Africa pay", "Pesapal, mobile money", "Cards, region limits", "Cards, region limits", "Partner billing"],
        ["Tax", "You name it and set %", "Tax codes / setups", "Tax rates / GST", "Tax engines"],
        ["Onboarding", "Vellisys walks you in", "Self-serve + partners", "Self-serve + advisors", "Implementers"],
        ["Mailbox", "Company Hostinger/Titan", "Your own SMTP / none", "Your own SMTP", "Desktop / add-ins"],
        ["Best when", "You sell on paper", "You need US payroll", "You have an accountant", "You already run Sage"],
    ]
    table = s.shapes.add_table(1 + len(rows), 5, Inches(0.45), Inches(1.55), Inches(12.4), Inches(5.1)).table
    widths = [Inches(2.15), Inches(2.55), Inches(2.55), Inches(2.55), Inches(2.6)]
    for i, w in enumerate(widths):
        table.columns[i].width = w
    for c, htext in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = htext
        for p in cell.text_frame.paragraphs:
            p.alignment = PP_ALIGN.LEFT
            for run in p.runs:
                set_run(run, 12, True, WHITE)
        cell.fill.solid()
        cell.fill.fore_color.rgb = NAVY
    for r, row in enumerate(rows, start=1):
        for c, val in enumerate(row):
            cell = table.cell(r, c)
            cell.text = val
            for p in cell.text_frame.paragraphs:
                for run in p.runs:
                    set_run(run, 12, c <= 1, BLUE if c == 1 else INK)
            cell.fill.solid()
            cell.fill.fore_color.rgb = PAPER if r % 2 == 1 else WHITE
    add_box(
        s,
        Inches(0.55),
        Inches(6.75),
        Inches(12.2),
        Inches(0.4),
        "We do not replace a Big Four audit pack. We replace the lost quote, the unbranded invoice, and the spreadsheet that only one person understands.",
        12,
        False,
        MUTED,
    )

    # 5 Books
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.7), Inches(10), Inches(0.4), "The books", 28, True, NAVY)
    add_box(s, Inches(0.55), Inches(1.12), Inches(10), Inches(0.3), "Everything that used to live in five folders", 15, False, MUTED)
    feats = [
        "Quotations that convert to invoices without retyping lines, tax or currency.",
        "Invoices with due dates, part receipts, and balances that stay visible until they are cleared.",
        "Receipts that print RECEIVED and DUE, in the same stationery as the invoice.",
        "Expenses against suppliers, with categories that feed reports.",
        "Delivery notes: quantities out, no prices. Return notes when goods come back.",
        "Refunds in and out, linked when you can, so Profit & Loss stays honest.",
        "Debtors list with branded reminders from the company mailbox.",
        "Creditors list with a note or a letter when you need to write to a supplier.",
        "Headed letters and custom documents on the same paper as the books.",
        "Reports: income, collections, outstanding, tax due, aging, quote conversion, CSV export.",
        "Profit & Loss on Crest: other income and costs beside invoices and expenses.",
        "Edit a saved sheet. Void when it is dead. Number formats you control.",
    ]
    for i, f in enumerate(feats):
        x = Inches(0.5 + (i % 2) * 6.4)
        y = Inches(1.5 + (i // 2) * 0.9)
        card(s, x, y, Inches(6.2), Inches(0.82))
        add_box(s, x + Inches(0.22), y + Inches(0.12), Inches(5.85), Inches(0.6), f, 13, False, INK)

    # 6 Desk
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.4), "The desk around the books", 28, True, NAVY)
    desk = [
        ("BRAND", "Primary and accent colours. Logo with optional white plate. Twelve layouts: Folio bar, Colour ledger, Corner bill, Accent bill, Twin copy, Accent stripe, Estate panel, Harbour block, watermarks and more. One choice reprints every sheet."),
        ("PEOPLE", "Up to three seats. Company admin opens Settings, Reports and the people list. Extra seats are Books or Sales so a salesperson cannot open the whole ledger."),
        ("MAIL", "Vellisys assigns a Hostinger or Titan mailbox. Quotes, invoices, receipts, letters and debtor reminders leave from that address with the logo on a white band. Personal Gmail stays out of the books."),
        ("PLANNER", "On Ledger and Crest: notes, budget targets, calendar, notifications. Receive on the bell when money should be logged."),
        ("SHARE", "Print the A4 sheet. Save PDF. WhatsApp the link. Email the sheet. The client sees paper, not the desk. On a phone the page is still A4, scaled to fit."),
        ("PAY US", "Pesapal in the same tab: mobile money, cards, bank. Or register without paying. Or book a demo. Super admin still walks the company in."),
    ]
    y = Inches(1.25)
    for title, body in desk:
        add_box(s, Inches(0.55), y, Inches(12.2), Inches(0.28), title, 13, True, BLUE)
        add_box(s, Inches(0.55), y + Inches(0.26), Inches(12.2), Inches(0.55), body, 14, False, INK)
        y += Inches(0.92)

    # 7 Packages
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.68), Inches(12), Inches(0.4), "Packages, billed per year", 28, True, NAVY)
    add_box(s, Inches(0.55), Inches(1.1), Inches(12), Inches(0.35), "List prices in Uganda shillings. Change currency on www.vellisys.com. Discount shown as was / now.", 14, False, MUTED)
    pkgs = [
        ("QUILL", "Starting", "UGX 150,000", "was 200,000  ·  1 seat  ·  per year", False, [
            "Company admin login",
            "Branded quotations, invoices, receipts",
            "Clients, debtors, email and WhatsApp share",
            "Print and PDF",
            "Reports for the person who signs in",
        ]),
        ("LEDGER", "Most companies", "UGX 200,000", "was 280,000  ·  2 seats  ·  per year", True, [
            "Admin plus one login (Books or Sales)",
            "Everything in Quill",
            "Expenses, creditors, delivery notes",
            "Planner notes, budget, calendar",
            "Headed letters from the company mailbox",
        ]),
        ("CREST", "Full house", "UGX 250,000", "was 350,000  ·  3 seats  ·  per year", False, [
            "Admin plus two logins",
            "Everything in Ledger",
            "Custom documents and every layout",
            "Profit & Loss, refunds, return notes",
            "Priority onboarding from Vellisys",
        ]),
    ]
    for i, (name, kicker, price, meta, hot, points) in enumerate(pkgs):
        x = Inches(0.45 + i * 4.25)
        y = Inches(1.55)
        bg = NAVY if hot else PAPER
        add_rect(s, x, y, Inches(4.05), Inches(5.45), bg)
        tc = WHITE if hot else NAVY
        accent = GOLD if hot else BLUE
        bodyc = RGBColor(0xD5, 0xDC, 0xF0) if hot else INK
        add_box(s, x + Inches(0.22), y + Inches(0.22), Inches(3.6), Inches(0.3), kicker, 13, True, accent)
        add_box(s, x + Inches(0.22), y + Inches(0.52), Inches(3.6), Inches(0.45), name, 26, True, tc)
        add_box(s, x + Inches(0.22), y + Inches(1.05), Inches(3.6), Inches(0.35), price, 18, True, accent)
        add_box(s, x + Inches(0.22), y + Inches(1.4), Inches(3.6), Inches(0.4), meta, 12, False, bodyc)
        add_lines(s, x + Inches(0.22), y + Inches(1.9), Inches(3.6), Inches(3.2), ["•  " + p for p in points], 14, bodyc)

    # 8 Ease
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.7), Inches(12), Inches(0.4), "How Vellisys eases the work", 28, True, NAVY)
    ease = [
        "Monday: raise a quote in the company colours. Share the link. When they accept, press Invoice. Lines, tax and currency copy.",
        "When money lands: record a receipt against the invoice. Part payment is a receipt, not a rewritten bill. Debtors drop by themselves.",
        "When goods leave: a delivery note with quantities. The invoice stays the bill. The driver carries paper that matches the office.",
        "When you spend: log the supplier and the tax. Reports tot income, expenses and tax due for the period you pick.",
        "When someone is late: open Debtors, send a reminder from the company mailbox. It looks like the rest of your stationery.",
        "When you hire a second person: give them Sales so they can quote without opening Settings or the full reports.",
        "When you open in another country: set currency, tax name and tax percent. You do not wait for a global vendor to add your rate.",
        "When the year turns: one invoice from Vellisys, not a surprise monthly SaaS that grew with seats you forgot.",
    ]
    y = Inches(1.25)
    for item in ease:
        add_rect(s, Inches(0.55), y + Inches(0.08), Inches(0.12), Inches(0.12), BLUE)
        add_box(s, Inches(0.85), y, Inches(11.8), Inches(0.65), item, 15, False, INK)
        y += Inches(0.7)

    # 9 Onboarding
    s = light_slide()
    add_box(s, Inches(0.55), Inches(0.68), Inches(12), Inches(0.4), "A note as you start", 28, True, NAVY)
    letter = [
        ("Welcome to Vellisys.", True),
        ("You are not buying another login to a foreign accounts cloud. You are opening a desk that prints like your office and keeps the books in one place.", False),
        ("This is how we start, together:", False),
        ("1. You pick Quill, Ledger or Crest, or you ask us to recommend one from how many people will sign in and whether you need Planner and Profit & Loss.", False),
        ("2. You pay on Pesapal, or we invoice you, or you register and we call. There is no password until the desk is opened on purpose.", False),
        ("3. We create the company, assign the sending mailbox, and send a welcome from info@vellisys.com with a short tutorial.", False),
        ("4. You sign in. Settings opens first: logo, colours, TIN, tax name and rate, currency, bank, document prefix. Save once. Every sheet follows.", False),
        ("5. Add the people who will work the desk. Add the first clients. Raise the first quotation. Convert it when they say yes.", False),
        ("We stay on the line for branding, the first documents, and the mailbox. If something is unclear, write to info@vellisys.com or call the numbers on the last page. You should never have to invent a workaround in Excel to make Vellisys look true.", False),
        ("We are glad you are here.", False),
        ("The Vellisys desk  ·  Front Star Digital", True),
    ]
    box = s.shapes.add_textbox(Inches(0.55), Inches(1.18), Inches(12.2), Inches(5.9))
    tf = box.text_frame
    tf.word_wrap = True
    for i, (line, strong) in enumerate(letter):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        run = p.add_run()
        run.text = line
        set_run(run, 16 if strong else 14, strong, NAVY if strong else INK)

    # 10 Close
    s = dark_slide()
    dark_rails(s)
    add_box(s, Inches(0.7), Inches(1.45), Inches(10), Inches(0.35), "Next step", 16, True, GOLD)
    add_box(s, Inches(0.7), Inches(1.9), Inches(11), Inches(0.7), "Open the desk with us.", 32, True, WHITE)
    add_box(
        s,
        Inches(0.7),
        Inches(2.75),
        Inches(10.5),
        Inches(1.4),
        "Pay at www.vellisys.com, register without paying, or book a demo. Tell us the company name, how many people will sign in, the currency you bill in, and the tax you charge. We will match a package and open the books.",
        16,
        False,
        MIST,
    )
    add_box(s, Inches(0.7), Inches(4.3), Inches(10), Inches(0.35), "www.vellisys.com", 20, True, GOLD)
    add_box(s, Inches(0.7), Inches(4.7), Inches(10), Inches(0.35), "info@vellisys.com", 20, True, WHITE)
    add_box(s, Inches(0.7), Inches(5.1), Inches(10), Inches(0.35), "+256 779 971 024", 20, True, WHITE)
    add_box(s, Inches(0.7), Inches(5.5), Inches(10), Inches(0.35), "+256 756 524 451", 20, True, WHITE)
    add_box(s, Inches(0.7), Inches(6.05), Inches(12), Inches(0.3), "Front Star Digital  ·  Vellisys branded books", 14, False, MIST)
    add_box(s, Inches(0.7), Inches(6.4), Inches(12), Inches(0.55), "This deck is for the company named in your email or meeting. Figures are yearly list prices in UGX as published on the site.\nThank you for reading. We will make the paper look like you.", 13, False, WHITE)

    for dest in dests:
        os.makedirs(os.path.dirname(dest), exist_ok=True)
        prs.save(dest)
        print(dest, os.path.getsize(dest), "bytes")


def main() -> int:
    asset_dir = prepare_assets()
    dests = [
        "/home/ubuntu/Desktop/Vellisys_Client_Pitch_Deck.pptx",
        "/opt/cursor/artifacts/Vellisys_Client_Pitch_Deck.pptx",
    ]
    build(asset_dir, dests)
    return 0


if __name__ == "__main__":
    sys.exit(main())
